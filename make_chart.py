import csv
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

wb = Workbook()
ws = wb.active
ws.title = "Cache Results"


with open("cache_results.csv", newline="") as file:
    reader = csv.reader(file)

    for i, row in enumerate(reader):

        if i == 0:
            # header
            ws.append(row)
        else:
            # convert everything to correct types
            row[0] = int(row[0])      # Run
            row[1] = int(row[1])      # Block Size
            row[2] = int(row[2])      # Associativity
            row[3] = int(row[3])      # Access Pattern
            row[4] = int(row[4])      # Stride
            row[5] = int(row[5])      # Hits
            row[6] = int(row[6])      # Misses
            row[7] = float(row[7])    # Hit Rate

            ws.append(row)

max_row = ws.max_row

data = Reference(ws, min_col=8, min_row=2, max_row=max_row)
categories = Reference(ws, min_col=1, min_row=2, max_row=max_row)


chart = BarChart()
chart.type = "col"
chart.grouping = "clustered"
chart.overlap = 0

chart.title = "Cache Hit Rate"
chart.title.overlay = False
chart.x_axis.title = "Run"
chart.x_axis.title_overlay = False
chart.y_axis.title = "Hit %"

chart.add_data(data, titles_from_data=False)
chart.set_categories(categories)

#Styling
chart.legend = None
chart.y_axis.majorGridlines = None

chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 100

chart.gapWidth = 60

chart.dataLabels = DataLabelList()
chart.dataLabels.showVal = True
chart.dataLabels.showSerName = False
chart.dataLabels.showCatName = False
chart.dataLabels.showLegendKey = False

chart.x_axis.delete = False


ws.add_chart(chart, "J2")

wb.save("Results.xlsx")